#!/usr/bin/env python3
"""Build a concise, editable slide deck for the staleness-adaptive proposal."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "STORM/STORM/crossover_experiment_qwen35_leakage_safe_k16_20260909"
OUTPUT = ROOT / "staleness-adaptive-communication-proposal_slides.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

BG = "F6F3EC"
PAPER = "FFFEFA"
INK = "14293B"
MUTED = "65727D"
FAINT = "DDE3E5"
BLUE = "2F6F91"
BLUE_LIGHT = "DDEBF2"
TEAL = "2A9D8F"
TEAL_LIGHT = "DDF1EC"
CORAL = "E76F51"
CORAL_LIGHT = "F9E3DC"
GOLD = "E3B64A"
GOLD_LIGHT = "F7EDD1"
GREEN = "4A9568"
GREEN_LIGHT = "E3F0E7"
RED = "B84C4C"
RED_LIGHT = "F5DEDE"
WHITE = "FFFFFF"
FONT = "Lato"
MONO = "Noto Sans Mono"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    color: str = INK,
    bold: bool = False,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    italic: bool = False,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    segments: list[tuple[str, bool, str]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, bold, color in segments:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_shape(
    slide,
    shape_type,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = PAPER,
    line: str | None = None,
    radius: bool = False,
):
    if radius:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: str = FAINT, width: float = 1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_pill(slide, text: str, x: float, y: float, w: float, *, fill: str, color: str, size: float = 10):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill=fill, radius=True)
    add_text(slide, text.upper(), x, y + 0.005, w, 0.31, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_card(slide, x: float, y: float, w: float, h: float, *, fill: str = PAPER, line: str = FAINT):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=True)


def add_header(slide, section: str, title: str, subtitle: str | None = None):
    add_text(slide, section.upper(), 0.72, 0.38, 5.7, 0.3, size=10, color=BLUE, bold=True)
    add_text(slide, title, 0.72, 0.76, 11.9, 0.58, size=27, color=INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.72, 1.36, 11.8, 0.45, size=13.5, color=MUTED)


def add_footer(slide, number: int, source: str = "Proposal v3 · corrected Qwen pilot · 10 Sep 2026"):
    add_line(slide, 0.72, 7.12, 12.62, 7.12, color=FAINT, width=0.75)
    add_text(slide, source, 0.72, 7.20, 10.8, 0.18, size=8.5, color=MUTED)
    add_text(slide, f"{number:02d}", 11.95, 7.18, 0.65, 0.2, size=9, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)


def add_bullet(slide, title: str, body: str, x: float, y: float, w: float, *, accent: str = TEAL):
    add_shape(slide, MSO_SHAPE.OVAL, x, y + 0.07, 0.12, 0.12, fill=accent)
    add_text(slide, title, x + 0.24, y, w - 0.24, 0.26, size=14, color=INK, bold=True)
    add_text(slide, body, x + 0.24, y + 0.29, w - 0.24, 0.55, size=11.5, color=MUTED, line_spacing=1.05)


def load_results():
    by_k = defaultdict(dict)
    with (RESULTS / "by_k_strategy.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            by_k[row["condition"]][int(row["k"])] = row

    overall = {}
    for condition, rows in by_k.items():
        values = list(rows.values())
        total_n = sum(int(row["n"]) for row in values)
        overall[condition] = {
            "success": sum(int(row["n"]) * float(row["recovery_success_rate"]) for row in values) / total_n,
            "tokens": sum(int(row["n"]) * float(row["mean_end_to_end_tokens"]) for row in values) / total_n,
        }

    bands = {}
    with (RESULTS / "robustness_contrasts.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            if row["comparator"] == "P3" and row["baseline"] == "P1" and row["stratum_type"] == "k_band":
                bands[row["stratum"]] = {
                    "effect": float(row["paired_success_difference"]),
                    "low": float(row["bootstrap_95_low"]),
                    "high": float(row["bootstrap_95_high"]),
                }

    policy = {}
    with (RESULTS / "policy_accuracy.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            policy[row["correct_action_group"]] = float(row["policy_action_accuracy"])

    p5 = []
    with (RESULTS / "replay_results.csv").open(newline="") as handle:
        for row in csv.DictReader(handle):
            if row["requested_condition"] == "P5":
                p5.append(row)
    route_success = defaultdict(list)
    for row in p5:
        bucket = "correct" if row["policy_action_correct"] == "True" else "wrong"
        route_success[bucket].append(row["recovery_success"] == "True")
    conditional = {key: sum(vals) / len(vals) for key, vals in route_success.items()}

    return overall, bands, policy, conditional


def slide_1(prs, overall):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_pill(slide, "Research proposal · corrected pilot", 0.75, 0.55, 2.75, fill=BLUE_LIGHT, color=BLUE)
    add_text(slide, "Staleness-adaptive\ncommunication", 0.75, 1.15, 6.2, 1.65, size=34, color=INK, bold=True)
    add_text(slide, "What should an agent receive when its write is rejected?", 0.78, 2.95, 5.65, 0.65, size=17, color=MUTED)

    add_card(slide, 7.15, 0.72, 5.36, 4.65, fill=PAPER, line=FAINT)
    add_text(slide, "ASYNC WRITE CONFLICT", 7.55, 1.08, 4.55, 0.3, size=10, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.OVAL, 7.65, 1.75, 1.12, 1.12, fill=BLUE_LIGHT)
    add_text(slide, "A", 7.65, 1.75, 1.12, 1.12, size=24, color=BLUE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_shape(slide, MSO_SHAPE.OVAL, 10.87, 1.75, 1.12, 1.12, fill=TEAL_LIGHT)
    add_text(slide, "B", 10.87, 1.75, 1.12, 1.12, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "read v₀", 7.51, 3.03, 1.4, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "lands v₁…vₖ", 10.65, 3.03, 1.55, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.82, 3.6, 2.05, 0.9, fill=INK, radius=True)
    add_text(slide, "shared.py\ncurrent = vₖ", 8.82, 3.63, 2.05, 0.82, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_shape(slide, MSO_SHAPE.DOWN_ARROW, 8.05, 3.55, 0.52, 0.82, fill=CORAL)
    add_text(slide, "WRITE\nREFUSED", 7.30, 4.52, 2.0, 0.55, size=11, color=CORAL, bold=True, align=PP_ALIGN.CENTER)

    stats = [("8", "scenario families"), ("16", "staleness levels"), ("3,456", "continuations")]
    for i, (value, label) in enumerate(stats):
        x = 0.75 + i * 2.08
        add_card(slide, x, 5.25, 1.83, 1.25, fill=PAPER, line=FAINT)
        add_text(slide, value, x + 0.12, 5.43, 1.59, 0.43, size=23, color=BLUE if i < 2 else TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.1, 5.95, 1.63, 0.28, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)

    add_rich_text(
        slide,
        [("Pilot verdict: ", True, INK), ("no deployable crossover yet", True, CORAL)],
        7.15,
        5.72,
        5.36,
        0.42,
        size=17,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "Natural, held-out confirmation is the next gate.", 7.25, 6.17, 5.15, 0.34, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 1)


def slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "01 · Problem", "A rejected write creates a communication decision", "The writer must recover from a world that changed after its last read.")

    y = 2.35
    nodes = [
        ("1", "Agent A reads", "shared.py · v₀", BLUE),
        ("k", "Other writes land", "v₁ … vₖ", GOLD),
        ("!", "A’s write is refused", "stale dependency", CORAL),
    ]
    for i, (tag, title, body, color) in enumerate(nodes):
        x = 0.8 + i * 2.05
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.68, 0.68, fill=color)
        add_text(slide, tag, x, y, 0.68, 0.68, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x - 0.18, y + 0.88, 1.55, 0.42, size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x - 0.18, y + 1.35, 1.55, 0.36, size=10, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 2:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 1.35, y + 0.17, 0.5, 0.34, fill=FAINT)

    add_text(slide, "What should the refusal say?", 0.8, 4.55, 5.55, 0.38, size=18, color=INK, bold=True)
    add_text(slide, "More information may help recovery—or amplify stale context.", 0.8, 5.0, 5.55, 0.58, size=13, color=MUTED)
    add_pill(slide, "Research question", 0.8, 5.75, 1.72, fill=CORAL_LIGHT, color=CORAL)
    add_text(slide, "Does the best payload depend on staleness?", 0.8, 6.18, 5.8, 0.46, size=17, color=CORAL, bold=True)

    cards = [
        ("P1 · STORM", "current content + diff\n+ stale dependencies", BLUE, BLUE_LIGHT),
        ("P3 · RATIONALE", "P1 + action-free\nedit reasoning", TEAL, TEAL_LIGHT),
        ("P5 · DIRECTIVE", "P1 + predicted route\n+ one-line hint", CORAL, CORAL_LIGHT),
    ]
    for i, (title, body, color, fill) in enumerate(cards):
        cy = 2.05 + i * 1.42
        add_card(slide, 7.0, cy, 5.5, 1.14, fill=PAPER, line=color)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, cy + 0.23, 1.65, 0.66, fill=fill, radius=True)
        add_text(slide, title, 7.2, cy + 0.23, 1.65, 0.66, size=11, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, body, 9.1, cy + 0.20, 3.0, 0.72, size=13, color=INK, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "Controls: P0 bare refusal · P1-pad token match · P4 full trajectory · P5-oracle upper bound", 7.02, 6.47, 5.45, 0.38, size=9.5, color=MUTED)
    add_footer(slide, 2)


def slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "02 · Proposal", "Isolate payload × staleness, then validate end-to-end", "Use replay for a large paired study; reserve live multi-agent runs for external validity.")

    steps = [
        ("1", "Harvest", "Natural STORM\nrefusals", BLUE, BLUE_LIGHT),
        ("2", "Snapshot", "Repo + agent\ncontext", GOLD, GOLD_LIGHT),
        ("3", "Replay", "Same episode,\nP0–P5", TEAL, TEAL_LIGHT),
        ("4", "Measure", "Mechanical\nrecovery", CORAL, CORAL_LIGHT),
        ("5", "Confirm", "Live runs on\n6 repos", GREEN, GREEN_LIGHT),
    ]
    for i, (num, title, body, color, fill) in enumerate(steps):
        x = 0.72 + i * 2.52
        add_card(slide, x, 2.17, 2.05, 1.92, fill=PAPER, line=color)
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.16, 2.34, 0.46, 0.46, fill=color)
        add_text(slide, num, x + 0.16, 2.34, 0.46, 0.46, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x + 0.74, 2.36, 1.12, 0.35, size=15, color=INK, bold=True)
        add_text(slide, body, x + 0.18, 3.00, 1.68, 0.74, size=13, color=MUTED, bold=False, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < 4:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.09, 2.88, 0.34, 0.44, fill=FAINT)

    add_text(slide, "Three views of staleness", 0.72, 4.62, 3.0, 0.4, size=16, color=INK, bold=True)
    measures = [
        ("TEMPORAL", "seconds since read", BLUE, BLUE_LIGHT),
        ("SEMANTIC", "changed code + symbols", TEAL, TEAL_LIGHT),
        ("INVESTMENT", "tokens + tool calls", GOLD, GOLD_LIGHT),
    ]
    for i, (name, desc, color, fill) in enumerate(measures):
        x = 0.72 + i * 2.43
        add_card(slide, x, 5.18, 2.18, 1.18, fill=fill, line=color)
        add_text(slide, name, x + 0.16, 5.38, 1.86, 0.26, size=10, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.16, 5.78, 1.86, 0.28, size=11, color=INK, align=PP_ALIGN.CENTER)

    add_card(slide, 8.25, 4.63, 4.35, 1.74, fill=INK, line=INK)
    add_text(slide, "Primary confirmatory test", 8.55, 4.91, 3.75, 0.3, size=11, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "payload × staleness", 8.5, 5.33, 3.85, 0.42, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "paired by episode · held-out threshold", 8.5, 5.87, 3.85, 0.28, size=10.5, color="C8D3DA", align=PP_ALIGN.CENTER)
    add_footer(slide, 3)


def slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "03 · Corrected pilot", "A leakage-safe mechanism study", "Local Qwen3.5-35B-A3B · controlled synthetic fixtures · episode-clustered inference")

    add_card(slide, 0.72, 2.0, 12.0, 1.28, fill=INK, line=INK)
    parts = [("8", "families"), ("16", "k levels"), ("9", "strategies"), ("3", "seeds")]
    for i, (value, label) in enumerate(parts):
        x = 1.1 + i * 2.52
        add_text(slide, value, x, 2.19, 1.2, 0.42, size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x, 2.69, 1.2, 0.27, size=10, color="C8D3DA", align=PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "×", x + 1.62, 2.29, 0.48, 0.38, size=20, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "=", 11.1, 2.27, 0.42, 0.4, size=22, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "3,456", 11.47, 2.19, 1.05, 0.42, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "continuations", 11.35, 2.69, 1.28, 0.27, size=10, color="C8D3DA", align=PP_ALIGN.CENTER)

    add_text(slide, "What changed after the invalid 100% P3 pilot", 0.72, 3.72, 6.05, 0.44, size=18, color=INK, bold=True)
    changes = [
        ("Action-free rationales", "No adapt / abandon / escalate cues."),
        ("Behavioral validators", "Code must work; comments cannot game scoring."),
        ("Protected write markers", "Every concurrent revision must survive."),
        ("Pre + post audits", "Exact prompts and saved outputs independently checked."),
    ]
    for i, (title, body) in enumerate(changes):
        col, row = i % 2, i // 2
        x, y = 0.72 + col * 3.05, 4.32 + row * 1.0
        add_card(slide, x, y, 2.82, 0.78, fill=PAPER, line=FAINT)
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.15, y + 0.18, 0.36, 0.36, fill=GREEN)
        add_text(slide, "✓", x + 0.15, y + 0.17, 0.36, 0.36, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x + 0.62, y + 0.12, 2.02, 0.24, size=11.5, color=INK, bold=True)
        add_text(slide, body, x + 0.62, y + 0.39, 2.02, 0.28, size=9.1, color=MUTED)

    add_card(slide, 7.25, 3.72, 5.47, 2.6, fill=GREEN_LIGHT, line=GREEN)
    add_text(slide, "AUDIT STATUS", 7.62, 4.08, 1.38, 0.3, size=10, color=GREEN, bold=True)
    add_text(slide, "PASS", 10.66, 3.98, 1.48, 0.46, size=24, color=GREEN, bold=True, align=PP_ALIGN.RIGHT)
    add_line(slide, 7.62, 4.6, 12.32, 4.6, color="BCD8C5", width=1)
    audit_rows = [
        ("P3 directive cues", "0"),
        ("Forbidden prompt fields", "0"),
        ("Rescore mismatches", "0 / 1,559"),
        ("Active validators", "8"),
    ]
    for i, (label, value) in enumerate(audit_rows):
        yy = 4.81 + i * 0.36
        add_text(slide, label, 7.62, yy, 3.05, 0.23, size=10.5, color=MUTED)
        add_text(slide, value, 10.77, yy, 1.53, 0.23, size=10.5, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    add_pill(slide, "Boundary", 0.72, 6.44, 1.05, fill=CORAL_LIGHT, color=CORAL)
    add_text(slide, "Synthetic mechanism evidence—not yet natural STORM/Commit0 or end-to-end pass rate.", 1.94, 6.46, 10.55, 0.29, size=11.3, color=MUTED)
    add_footer(slide, 4)


def slide_5(prs, overall):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "04 · Result", "More feedback was not monotonically better", "Recovery success across 384 continuations per strategy; oracle shown only as a diagnostic upper bound.")

    strategies = [
        ("P5-oracle", "P5 oracle", GREEN),
        ("P3", "P3 rationale", TEAL),
        ("P1", "P1 STORM", BLUE),
        ("P1-pad", "P1 padded", GOLD),
        ("P4", "P4 trajectory", CORAL),
        ("P5", "P5 predicted", RED),
    ]
    chart_x, chart_y, chart_w = 2.15, 2.08, 6.05
    for i, (key, label, color) in enumerate(strategies):
        yy = chart_y + i * 0.71
        value = overall[key]["success"] * 100
        add_text(slide, label, 0.75, yy + 0.08, 1.25, 0.28, size=11, color=INK, bold=key in {"P1", "P3"}, align=PP_ALIGN.RIGHT)
        add_shape(slide, MSO_SHAPE.RECTANGLE, chart_x, yy + 0.09, chart_w, 0.26, fill="E5E8E7")
        add_shape(slide, MSO_SHAPE.RECTANGLE, chart_x, yy + 0.09, chart_w * value / 100, 0.26, fill=color)
        add_text(slide, f"{value:.1f}%", chart_x + chart_w + 0.12, yy, 0.72, 0.42, size=12, color=color, bold=True)
    add_text(slide, "0", chart_x - 0.04, 6.43, 0.25, 0.23, size=8.5, color=MUTED)
    add_text(slide, "50", chart_x + chart_w * 0.5 - 0.12, 6.43, 0.3, 0.23, size=8.5, color=MUTED)
    add_text(slide, "100%", chart_x + chart_w - 0.4, 6.43, 0.42, 0.23, size=8.5, color=MUTED)

    add_card(slide, 9.08, 2.02, 3.55, 1.55, fill=TEAL_LIGHT, line=TEAL)
    add_text(slide, "P3 − P1", 9.4, 2.29, 2.9, 0.28, size=11, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "+1.6 pts", 9.32, 2.67, 3.05, 0.5, size=25, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "95% CI −2.6 to +5.7 · p = .510", 9.29, 3.2, 3.14, 0.25, size=9.6, color=MUTED, align=PP_ALIGN.CENTER)

    add_card(slide, 9.08, 3.84, 3.55, 1.28, fill=CORAL_LIGHT, line=CORAL)
    add_text(slide, "P4 − P1", 9.4, 4.09, 2.9, 0.26, size=10.5, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "−12.0 pts", 9.35, 4.45, 3.0, 0.4, size=22, color=CORAL, bold=True, align=PP_ALIGN.CENTER)

    add_card(slide, 9.08, 5.39, 3.55, 1.22, fill=PAPER, line=FAINT)
    add_text(slide, "Interpretation", 9.36, 5.61, 2.95, 0.24, size=10.5, color=BLUE, bold=True)
    add_text(slide, "Rationale ≈ baseline.\nFull trajectory actively hurts.", 9.36, 5.93, 2.95, 0.52, size=13, color=INK, bold=True)
    add_footer(slide, 5)


def slide_6(prs, bands):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "05 · Staleness", "The observed pattern reverses at higher k", "P3 minus P1 recovery success by staleness band; whiskers are episode-clustered 95% bootstrap intervals.")

    chart_x, chart_y, chart_w, chart_h = 0.95, 2.18, 7.1, 3.8
    y_min, y_max = -15, 22

    def sy(value):
        return chart_y + chart_h * (y_max - value) / (y_max - y_min)

    for tick in [-10, 0, 10, 20]:
        yy = sy(tick)
        add_line(slide, chart_x, yy, chart_x + chart_w, yy, color=INK if tick == 0 else FAINT, width=1.5 if tick == 0 else 0.7)
        add_text(slide, f"{tick:+d}" if tick else "0", 0.42, yy - 0.12, 0.42, 0.24, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "percentage points", 0.18, 1.95, 1.25, 0.25, size=8.7, color=MUTED)

    names = ["k1-4", "k5-8", "k9-12", "k13-16"]
    for i, name in enumerate(names):
        item = bands[name]
        effect, low, high = item["effect"] * 100, item["low"] * 100, item["high"] * 100
        cx = chart_x + 0.95 + i * 1.67
        zero_y, effect_y = sy(0), sy(effect)
        top, height = min(zero_y, effect_y), abs(zero_y - effect_y)
        color = TEAL if effect >= 0 else CORAL
        add_shape(slide, MSO_SHAPE.RECTANGLE, cx - 0.28, top, 0.56, max(height, 0.02), fill=color)
        add_line(slide, cx, sy(low), cx, sy(high), color=INK, width=1.25)
        add_line(slide, cx - 0.12, sy(low), cx + 0.12, sy(low), color=INK, width=1.25)
        add_line(slide, cx - 0.12, sy(high), cx + 0.12, sy(high), color=INK, width=1.25)
        label_y = effect_y - 0.4 if effect >= 0 else effect_y + 0.08
        add_text(slide, f"{effect:+.1f}", cx - 0.42, label_y, 0.84, 0.3, size=11, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, name, cx - 0.48, 6.15, 0.96, 0.29, size=10.5, color=INK, bold=True, align=PP_ALIGN.CENTER)

    add_card(slide, 8.65, 2.17, 3.95, 1.45, fill=CORAL_LIGHT, line=CORAL)
    add_text(slide, "Original H2", 8.97, 2.43, 3.28, 0.28, size=10.5, color=CORAL, bold=True)
    add_text(slide, "Rich feedback should help\nmore as staleness grows.", 8.97, 2.81, 3.28, 0.6, size=15, color=INK, bold=True)

    add_card(slide, 8.65, 3.92, 3.95, 1.9, fill=PAPER, line=FAINT)
    add_text(slide, "Pilot evidence", 8.97, 4.19, 3.28, 0.28, size=10.5, color=BLUE, bold=True)
    add_text(slide, "No consistently positive crossover.", 8.97, 4.58, 3.25, 0.43, size=15, color=INK, bold=True)
    add_text(slide, "The sign flips after k=8, but action mix also changes with k.", 8.97, 5.11, 3.2, 0.54, size=11.3, color=MUTED)

    add_pill(slide, "Read cautiously", 8.65, 6.16, 1.55, fill=GOLD_LIGHT, color="946F13")
    add_text(slide, "Descriptive target for a natural, action-balanced study—not a causal threshold.", 10.33, 6.14, 2.27, 0.54, size=9.2, color=MUTED)
    add_footer(slide, 6)


def slide_7(prs, overall, policy, conditional):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "06 · Mechanism", "Direction works—route prediction does not", "P5 succeeds when the router is right, but adapt cases are mostly misrouted.")

    add_card(slide, 0.72, 2.08, 7.25, 3.98, fill=PAPER, line=FAINT)
    add_text(slide, "P5 routing path", 1.03, 2.4, 2.0, 0.3, size=11, color=BLUE, bold=True)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.04, 3.07, 1.72, 0.92, fill=BLUE_LIGHT, radius=True)
    add_text(slide, "Observable\nepisode state", 1.04, 3.08, 1.72, 0.9, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 2.98, 3.34, 0.72, 0.4, fill=FAINT)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 3.91, 2.9, 1.7, 1.26, fill=INK, radius=True)
    add_text(slide, "ROUTER\n58.6%", 3.91, 2.91, 1.7, 1.24, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 5.84, 3.34, 0.72, 0.4, fill=FAINT)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.78, 2.51, 0.88, 0.88, fill=GREEN_LIGHT, line=GREEN, radius=True)
    add_text(slide, "✓", 6.78, 2.51, 0.88, 0.88, size=22, color=GREEN, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, f"{conditional['correct'] * 100:.1f}% success\nif route correct", 5.88, 2.05, 1.75, 0.52, size=10.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.78, 4.13, 0.88, 0.88, fill=RED_LIGHT, line=RED, radius=True)
    add_text(slide, "×", 6.78, 4.13, 0.88, 0.88, size=22, color=RED, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, f"{conditional['wrong'] * 100:.1f}% success\nif route wrong", 5.88, 5.08, 1.75, 0.52, size=10.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 5.58, 3.52, 6.5, 2.97, color=GREEN, width=1.4)
    add_line(slide, 5.58, 3.55, 6.5, 4.57, color=RED, width=1.4)

    add_text(slide, "Router accuracy by action", 1.04, 4.66, 2.5, 0.28, size=10.5, color=MUTED, bold=True)
    actions = [("adapt", TEAL), ("abandon", BLUE), ("escalate", GOLD)]
    for i, (name, color) in enumerate(actions):
        x = 1.04 + i * 1.48
        value = policy[name] * 100
        add_text(slide, name, x, 5.03, 1.24, 0.22, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
        add_text(slide, f"{value:.1f}%", x, 5.31, 1.24, 0.35, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)

    add_card(slide, 8.4, 2.08, 4.22, 3.98, fill=INK, line=INK)
    add_text(slide, "DEPLOYMENT GAP", 8.75, 2.42, 3.55, 0.28, size=10, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{overall['P5']['success'] * 100:.1f}%", 8.72, 2.93, 1.55, 0.56, size=27, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "P5 predicted", 8.72, 3.5, 1.55, 0.28, size=9.7, color="C8D3DA", align=PP_ALIGN.CENTER)
    add_text(slide, "vs", 10.28, 3.1, 0.42, 0.28, size=11, color="C8D3DA", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{overall['P5-oracle']['success'] * 100:.1f}%", 10.71, 2.93, 1.55, 0.56, size=27, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "P5 oracle", 10.71, 3.5, 1.55, 0.28, size=9.7, color="C8D3DA", align=PP_ALIGN.CENTER)
    add_line(slide, 8.85, 4.02, 12.12, 4.02, color="3D5364", width=1)
    add_text(slide, "≈84.5%", 8.77, 4.34, 3.42, 0.47, size=23, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "router accuracy needed to match P3*", 8.77, 4.83, 3.42, 0.27, size=9.7, color="C8D3DA", align=PP_ALIGN.CENTER)
    add_text(slide, "1,344 tokens P5  vs  766 P3", 8.77, 5.39, 3.42, 0.3, size=11.5, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "*If conditional success rates remain fixed.", 8.78, 5.78, 3.4, 0.22, size=8.4, color="AAB8C1", align=PP_ALIGN.CENTER)
    add_footer(slide, 7)


def slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H, fill=BG)
    add_header(slide, "07 · Decision", "Do not deploy the adaptive protocol from this pilot", "The result changes the next experiment—not the motivation for the research question.")

    add_card(slide, 0.72, 2.06, 3.55, 3.85, fill=CORAL_LIGHT, line=CORAL)
    add_pill(slide, "Now", 1.03, 2.39, 0.75, fill=CORAL, color=WHITE)
    add_text(slide, "Keep P1 as the safe baseline", 1.03, 2.98, 2.93, 0.68, size=20, color=INK, bold=True)
    add_text(slide, "No reliable P3 advantage.\nP4 hurts. P5 routing is not ready.", 1.03, 3.87, 2.93, 0.84, size=13.5, color=MUTED)
    add_text(slide, "Adaptive k=4: 69.0%\nP1 baseline: 81.0%", 1.03, 5.05, 2.93, 0.58, size=13, color=CORAL, bold=True)

    add_text(slide, "Next study", 4.88, 2.07, 2.2, 0.38, size=18, color=INK, bold=True)
    next_steps = [
        ("01", "Natural refusals", "Harvest STORM/Commit0 episodes with natural assignments."),
        ("02", "Action-balanced + held out", "Pre-register payload × staleness; freeze thresholds and router."),
        ("03", "Replicate + validate", "Second model family, then end-to-end runs on six repos."),
    ]
    for i, (num, title, body) in enumerate(next_steps):
        y = 2.66 + i * 1.06
        add_shape(slide, MSO_SHAPE.OVAL, 4.88, y, 0.55, 0.55, fill=TEAL if i < 2 else BLUE)
        add_text(slide, num, 4.88, y, 0.55, 0.55, size=9.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, 5.65, y - 0.01, 2.5, 0.3, size=13, color=INK, bold=True)
        add_text(slide, body, 5.65, y + 0.34, 3.0, 0.48, size=10.5, color=MUTED)

    add_card(slide, 9.15, 2.07, 3.46, 3.85, fill=INK, line=INK)
    add_text(slide, "WHAT WOULD CHANGE\nTHE DECISION?", 9.48, 2.42, 2.8, 0.6, size=11, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "A held-out, natural\npayload × staleness\ninteraction", 9.48, 3.24, 2.8, 1.28, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_line(slide, 9.54, 4.75, 12.21, 4.75, color="3D5364", width=1)
    add_text(slide, "plus a router that clears\nthe pre-registered target", 9.48, 5.05, 2.8, 0.55, size=11.5, color="C8D3DA", align=PP_ALIGN.CENTER)

    add_rich_text(
        slide,
        [("Takeaway: ", True, BLUE), ("the right question is not “more or less context?” but “which context, for this conflict, now?”", True, INK)],
        0.78,
        6.37,
        11.85,
        0.48,
        size=15,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 8)


def set_metadata(prs):
    props = prs.core_properties
    props.title = "Staleness-Adaptive Communication in Asynchronous Multi-Agent Code Generation"
    props.subject = "Proposal overview and leakage-safe Qwen pilot results"
    props.author = "OpenAI Codex"
    props.keywords = "multi-agent, code generation, staleness, communication, STORM, Qwen"
    props.comments = "Generated from the proposal v3 and the 2026-09-09 corrected experiment artifacts."


def main():
    overall, bands, policy, conditional = load_results()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    set_metadata(prs)
    slide_1(prs, overall)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs, overall)
    slide_6(prs, bands)
    slide_7(prs, overall, policy, conditional)
    slide_8(prs)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
